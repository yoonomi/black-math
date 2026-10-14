# -*- coding: utf-8 -*-
"""按《创新实践》指导书附件2生成课程设计报告 Word。"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = Path(__file__).resolve().parent / "金科丞+影神图多智能体问答-课程设计报告.docx"


def set_run_font(run, size=10.5, bold=False, east="宋体", ascii_font="Times New Roman"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = ascii_font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), east)


def set_paragraph_format(p, first_indent=True, align="left", space_after=6):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(0)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Cm(0)
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pf.first_line_indent = Cm(0)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = Cm(0.74) if first_indent else Cm(0)


def add_text(p, text, size=10.5, bold=False, east="宋体"):
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, east=east)
    return run


def title(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_indent=False, align="center", space_after=12)
    add_text(p, text, size=16, bold=True)


def h1(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_indent=False, align="left", space_after=8)
    add_text(p, text, size=12, bold=True)


def h2(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_indent=False, align="left", space_after=6)
    add_text(p, text, size=10.5, bold=True)


def body(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_indent=True)
    add_text(p, text, size=10.5)
    return p


def caption_fig(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_indent=False, align="center", space_after=10)
    add_text(p, text, size=10.5, bold=True)


def caption_tab(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_indent=False, align="center", space_after=4)
    add_text(p, text, size=10.5, bold=True)


def pre(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_indent=False, space_after=6)
    pf = p.paragraph_format
    pf.line_spacing = 1.0
    run = p.add_run(text)
    set_run_font(run, size=9, east="宋体", ascii_font="Consolas")


def set_cell_font(cell, text, bold=False, size=9, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    add_text(p, text, size=size, bold=bold)


def shade_header(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_three_line(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    specs = [
        ("top", "12"),
        ("left", "0"),
        ("bottom", "12"),
        ("right", "0"),
        ("insideH", "0"),
        ("insideV", "0"),
    ]
    for edge, sz in specs:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil" if sz == "0" else "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)
    for cell in table.rows[0].cells:
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcBorders = OxmlElement("w:tcBorders")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "0")
        bottom.set(qn("w:color"), "000000")
        tcBorders.append(bottom)
        tcPr.append(tcBorders)
        shade_header(cell)


def set_box_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tblPr.append(borders)


def add_flow_table(doc, headers, rows):
    """用表格代替等宽字符框图，避免 Word 中错位。"""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        set_cell_font(table.rows[0].cells[i], h, bold=True, size=10.5, center=True)
        shade_header(table.rows[0].cells[i])
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            center = c == 0
            set_cell_font(table.rows[r + 1].cells[c], str(val), size=9, center=center)
    set_box_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        set_cell_font(table.rows[0].cells[i], h, bold=True, size=10.5, center=True)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            set_cell_font(table.rows[r + 1].cells[c], str(val), size=9, center=True)
    set_three_line(table)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    return table


def formula(doc, text):
    p = doc.add_paragraph()
    set_paragraph_format(p, first_indent=False, align="center", space_after=8)
    add_text(p, text, size=10.5, east="Times New Roman")


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.6)
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    title(doc, "基于 GraphRAG 的《黑神话：悟空》影神图多智能体问答系统")
    p = doc.add_paragraph()
    set_paragraph_format(p, first_indent=False, align="center", space_after=4)
    add_text(p, "课程设计报告", size=12, bold=True)

    p = doc.add_paragraph()
    set_paragraph_format(p, first_indent=False, align="center", space_after=4)
    add_text(p, "学号：2023337621104　　姓名：金科丞", size=10.5)

    p = doc.add_paragraph()
    set_paragraph_format(p, first_indent=False, align="center", space_after=4)
    add_text(p, "学号：　　　　　　　　姓名：　　　　（组员二，提交前填写）", size=10.5)

    p = doc.add_paragraph()
    set_paragraph_format(p, first_indent=False, align="center", space_after=12)
    add_text(p, "浙江理工大学信息学院　《创新实践》", size=10.5)

    h1(doc, "1．引言")
    h2(doc, "1.1 题目简介")
    body(
        doc,
        "本课题面向浙江理工大学信息学院《创新实践》课程，完成一套可运行的影神图问答与评测系统。系统以开源项目 Chat-BlackMyth 的 TinyGraph 为数据层，将《黑神话：悟空》影神图文本建成 Neo4j 知识图谱，在查询期引入规划、双路检索与校验三个角色协同作答，并在同一张图上对照智谱、DeepSeek、通义千问的图谱问答与纯生成表现。源码仓库为 https://github.com/yoonomi/black-math。",
    )

    h2(doc, "1.2 背景与意义")
    body(
        doc,
        "影神图是游戏内的人物、头目与关系说明，结构接近知识条目，但演示与问答仍多用自然语言。直接调用大模型时，容易把《西游记》原著常识答成游戏设定，例如将「黑风山土地」讲成小说情节。向量检索只能召回相似段落，对「部下、交手、位于」这类多跳关系不够稳定。",
    )
    body(
        doc,
        "知识图谱把「谁—关系—谁」显式存下来，图上行走可以补全跨条目的关系。微软 GraphRAG 与开源 TinyGraph 已验证「文本→图→检索生成」可行。Chat-BlackMyth 进一步把影神图做成可演示的 GraphRAG，但查询仍是一次局部检索，缺少可解释的多步推理，也缺少跨模型的封闭评测。本课题在其数据层上补齐查询智能体、公平对照与性能看板，使课程设计同时具备可运行系统、可解释路径和可讨论的实验结论。",
    )

    h2(doc, "1.3 目标")
    body(doc, "（1）在 Neo4j 中建成前三章（黑风山、黄风岭、小西天）影神图，实体带向量，关系可按模板走 1～3 步。")
    body(doc, "（2）查询期按「规划 → 双路检索 → 成答 → 校验」协作，界面展示子问题、实体、三元组与检索子图。")
    body(doc, "（3）图谱只建一次、向量固定智谱，问答切换智谱、DeepSeek、千问，对照「图谱问答」与「纯生成」。")
    body(doc, "（4）用封闭题集给出正确率、引用覆盖、时延与综合分，并总结「基线越弱、接图谱提升越大」。")

    h1(doc, "2．系统详细设计")
    h2(doc, "2.1 总体目标与模块划分")
    body(doc, "系统分为交互层、智能体层、图谱层与评测层。自上而下的调用关系如图 1 所示。")
    add_flow_table(
        doc,
        ["层次", "组成", "数据流"],
        [
            [
                "交互层",
                "Gradio 网页",
                "问答页选择模型后，图谱问答与同模型纯生成并行；性能页读取评测结果，展示综合排名、雷达图与分项图",
            ],
            ["模型层", "模型工厂", "智谱、DeepSeek、千问，统一聊天接口"],
            [
                "智能体层",
                "查询循环",
                "规划器 → 双路检索 → 成答器 → 校验器，最多 4 轮",
            ],
            [
                "图谱层",
                "TinyGraph",
                "原文邻域检索，以及按实体名的 1～3 步模板图查询",
            ],
            [
                "存储层",
                "Neo4j",
                "实体与关系；建图向量固定为智谱 embedding-3",
            ],
        ],
    )
    caption_fig(doc, "图1 系统模块结构")
    body(doc, "各模块职责见表 1。")
    caption_tab(doc, "表1 功能模块与职责")
    add_table(
        doc,
        ["模块", "主要文件", "职责"],
        [
            ["语料", "download_data/、portraits_ch1_3.md", "抓取影神图；筛前三章 90 人"],
            ["建图", "graphrag/graph.py", "分块、抽实体与三元组、写入 Neo4j、向量化"],
            ["模型工厂", "graphrag/llm/factory.py 等", "按名称创建聊天模型，统一 predict()"],
            ["规划器", "agents/planner.py", "单跳直接答；多跳按缺口追加子问题"],
            ["双路检索", "agents/retrieval.py", "关系走模板图查询；生平取原文"],
            ["成答与校验", "agents/loop.py、verifier.py", "通顺中文作答；证据不足则再检索"],
            ["交互", "webdemo.py、agents/vis.py", "对照对话、检索子图、综合性能图"],
            ["评测", "eval/run_eval.py", "14 题 × 三模型 × 两种模式"],
        ],
    )

    h2(doc, "2.2 建图模块：算法与流程")
    body(
        doc,
        "建图沿用 TinyGraph，按角色切块，避免固定字数滑窗把同一人物拆碎。处理流程如图 2 所示。",
    )
    add_flow_table(
        doc,
        ["步骤", "处理内容"],
        [
            ["1", "读入影神图 Markdown"],
            ["2", "按「影神图角色」切块，一人一块"],
            ["3", "对每个新块抽取实体与三元组"],
            ["4", "按同名合并，写入 Neo4j"],
            ["5", "社区发现（有插件则 Leiden，否则单社区回退）"],
            ["6", "实体描述向量化（embedding-3）"],
            ["7", "标记文档已加载，支持增量补块"],
        ],
    )
    caption_fig(doc, "图2 建图处理流程")
    body(
        doc,
        "实体节点主要属性包括 name、description、entity_id、embedding、communityIds。关系边属性为 name（关系名）。块与文档记录在本地 workspace/，用内容哈希去重，避免重复调用大模型。",
    )
    body(
        doc,
        "图查询不由模型生成任意语句，只使用绑定实体名的模板：从名称包含给定字符串的实体出发，沿无向关系走 1～3 步，返回头、关系、尾，从而保证模型不能随便写图数据库语句。",
    )

    h2(doc, "2.3 查询智能体：逻辑流程与数据结构")
    body(
        doc,
        "一次提问的控制流程如图 3 所示。这是相对原仓库的核心增量：原项目一次局部检索即生成；本系统把「要想什么、去哪查、答完核不核」拆成四个角色。",
    )
    add_flow_table(
        doc,
        ["步骤", "角色", "作用"],
        [
            ["1", "规划器", "判断直接作答，或追加子问题"],
            [
                "2",
                "双路检索",
                "关系类（部下、交手、位于等）走 1～3 步模板图查询；生平与身份取原文邻域，并始终拼接原文",
            ],
            ["3", "成答器", "只用证据写 1～2 句通顺中文"],
            ["4", "校验器", "证据足够则结束；不足则回到规划器，最多 4 轮"],
            ["5", "输出", "答案、子问题路径、实体与三元组，并绘制检索子图"],
        ],
    )
    caption_fig(doc, "图3 查询期智能体流程")
    body(doc, "共享证据集字段见表 2。")
    caption_tab(doc, "表2 证据集主要字段")
    add_table(
        doc,
        ["字段", "含义"],
        [
            ["entities", "命中实体（名、描述、编号）"],
            ["triples", "命中三元组（头、关系、尾）"],
            ["chunks", "原文摘录（截断，过滤「是 / 为」等弱关系）"],
            ["subquestions", "本轮子问题序列"],
            ["notes", "检索失败等备注"],
        ],
    )
    body(
        doc,
        "成答提示要求问什么答什么，禁止「证据提到」等过程套话，优先游戏设定而非原著课文。后处理去掉英文题型标签和「证据提到」之后的堆砌，减轻病句。校验器在完全没有实体、三元组与原文时才判为未知；有证据则根据草稿判断是否还需要追问。",
    )

    h2(doc, "2.4 模块输入输出")
    caption_tab(doc, "表3 主要模块输入输出")
    add_table(
        doc,
        ["模块", "输入", "输出"],
        [
            ["建图", "影神图路径、智谱密钥、Neo4j 地址", "图中实体与边、本地块缓存"],
            ["规划器", "问题、当前证据、轮次", "「直接作答」或「追加：子问题」"],
            ["双路检索", "子问题、图连接", "写入证据集"],
            ["成答器", "问题、证据文本", "中文答案"],
            ["校验器", "问题、草稿、证据", "通过 / 需补充 / 未知"],
            ["问答页", "用户问题、模型名", "左右答案、路径、子图"],
            ["评测", "题集 JSON", "results.json（逐题与汇总）"],
        ],
    )

    h2(doc, "2.5 用户界面")
    body(
        doc,
        "界面用 Gradio，分为两个标签页。（1）问答页：模型单选（智谱 / DeepSeek / 千问）；「连接图谱」默认只连接已有 Neo4j，勾选重建才会重新抽实体。左右两栏并行请求：左为智能体图谱问答，右为同模型纯生成。下方为推理过程、检索子图、实体表与三元组表。示例问题包括土地公、白衣秀士、广智广谋等稳题。",
    )
    body(
        doc,
        "（2）性能对照页：读取 eval/results.json。顶部自动生成综合结论；随后为三家图谱问答的综合排名、雷达画像、相对纯生成的 Hit@1 增益，以及分项柱状图与题型正确率。综合分定义为：",
    )
    formula(doc, "综合分 = 0.45×正确率 + 0.25×图谱引用 + 0.15×答案重叠 + 0.15×速度　　（1）")
    body(
        doc,
        "速度以三家图谱问答中最快者为 1。权重可调，分项均在表中列出，避免只看一个总分。",
    )

    h1(doc, "3．实验结果与讨论")
    h2(doc, "3.1 实验条件")
    body(doc, "图谱为前三章 90 个角色，Neo4j 中 441 个实体（均有向量）、910 条关系。")
    body(doc, "问答模型为智谱 glm-4-flash、DeepSeek deepseek-chat、千问 qwen3.5-flash（关闭思考链）。")
    body(doc, "建图与检索向量始终使用智谱 embedding-3。题集 14 题，含事实、多跳关系、对比；关键词与别名匹配计算 Hit@1。每家模型对照「图谱问答」与「纯生成」，共六路。")

    h2(doc, "3.2 结果展示")
    caption_tab(doc, "表4 六路评测汇总")
    add_table(
        doc,
        ["模型 / 模式", "Hit@1", "引用覆盖", "平均时延 / s"],
        [
            ["智谱 图谱问答", "0.64", "0.70", "4.3"],
            ["智谱 纯生成", "0.43", "0.88", "12.3"],
            ["DeepSeek 图谱问答", "0.79", "0.86", "3.9"],
            ["DeepSeek 纯生成", "0.71", "0.92", "2.7"],
            ["千问 图谱问答", "0.79", "0.93", "6.6"],
            ["千问 纯生成", "0.64", "0.87", "7.5"],
        ],
    )
    caption_tab(doc, "表5 三家图谱问答综合能力")
    add_table(
        doc,
        ["排名", "智能体", "综合分", "正确率", "图谱引用", "相对纯生成 Hit@1 提升"],
        [
            ["1", "DeepSeek", "0.76", "0.79", "0.86", "+0.07"],
            ["2", "千问", "0.72", "0.79", "0.93", "+0.14"],
            ["3", "智谱", "0.67", "0.64", "0.70", "+0.21"],
        ],
    )
    body(
        doc,
        "网页「性能对照」可加载与表 4、表 5 一致的柱状图、雷达图与综合结论文案。问答页在「黑风山土地是谁」一类题目上，左侧能约束到游戏土地公，右侧纯生成常讲原著；检索子图显示本轮实体与关系。",
    )

    h2(doc, "3.3 结果分析与讨论")
    body(
        doc,
        "（1）图谱对弱模型帮助更大。智谱纯生成 Hit@1 仅 0.43，接图谱后 0.64，提升 21 个百分点；千问提升 14 个百分点；最强基线 DeepSeek 从 0.71 到 0.79，仅 7 个百分点。弱模型更容易把原著常识当成游戏设定，图上的实体名和原文摘录起到约束作用。因此不能只用「DeepSeek 纯生成已经 0.71、智谱图谱问答只有 0.64」否定图谱：跨模型比的是综合能力，同一模型「有图 / 无图」才说明检索是否有效。",
    )
    body(
        doc,
        "（2）正确率天花板由较强聊天模型决定。图谱问答下 DeepSeek 与千问并列 0.79，智谱 0.64。综合分 DeepSeek 第一，因其速度与正确率都高；千问引用覆盖最高（0.93），答案更常点名影神图实体。",
    )
    body(
        doc,
        "（3）速度方面，左右栏已并行，页面等待接近较慢一侧。千问默认思考链曾导致单次约 30 秒、多轮到两分钟，关闭思考后图谱问答约 6.6 秒，与另外两家同一量级。DeepSeek 纯生成最快（2.7 秒）。",
    )
    body(
        doc,
        "（4）失败案例：「黄眉是谁、第几章、在哪」曾被答成「龟将部下、苦海、第二章」。设定上黄眉是第三章小西天妖王，据点小雷音寺，在苦海与龟蛇二将交手而非部下。图中弱边「黄眉—与—龟将」「龟将—位于—苦海」被模型脑补成隶属与章数。说明模板查询仍会带回「是 / 与 / 人物」一类空边。演示宜用土地公、白衣秀士、广智广谋；黄眉作为局限案例在报告中保留。",
    )
    body(
        doc,
        "（5）成答质量方面，早期存在「证据提到……」的分号堆砌。收紧成答提示并过滤过程套话后，土地归属类问题可稳定答成「黑风山土地管黑风山，黄风岭土地管黄风岭」，不再夹带灵吉菩萨等无关实体。这提升的是语句质量，封闭集 Hit@1 未因这一改动全部重跑。",
    )
    body(
        doc,
        "（6）题量为 14 题，偏少，结论表述为趋势而非精确到百分点的刷分。同时给出引用覆盖与时延，避免只看一个正确率。",
    )

    h1(doc, "4．结论")
    body(
        doc,
        "本课程设计在 Chat-BlackMyth / TinyGraph 数据层上，完成了可运行的影神图多智能体问答与跨模型评测。相对原开源项目，查询由一次局部生成改为规划、双路检索、成答、校验的协作；图跳转由程序模板约束；同一张图上比较三家大模型；并用封闭集得到「基线越弱、图谱提升越大」的实验结论。",
    )
    body(
        doc,
        "工作仍有限：演示库为前三章而非全量六章；评测 14 题；社区报告未作为主路径；弱关系边会带来如黄眉一类事实错误；Grok 未进入数字。后续可过滤空边、扩大封闭集、在同一流水线上补全量入库。",
    )
    body(
        doc,
        "对课程目标而言，小组完成了选题、可行性分析、系统实现、对照实验与文档，源码见 GitHub，满足任务书中「可运行程序、设计报告、演示」的交付要求。",
    )

    h1(doc, "参考文献")
    refs = [
        "[1] Edge D, Trinh H, Cheng N, et al. From local to global: A graph RAG approach to query-focused summarization[EB/OL]. https://arxiv.org/abs/2404.16130, 2024.",
        "[2] Nijikadesu. Chat-BlackMyth[CP/OL]. https://github.com/Nijikadesu/Chat-BlackMyth, 2024.",
        "[3] limafang. tiny-graphrag[CP/OL]. https://github.com/limafang/tiny-graphrag, 2024.",
        "[4] Neo4j, Inc. Neo4j Graph Database Documentation[EB/OL]. https://neo4j.com/docs/, 2025.",
        "[5] 智谱 AI. 智谱开放平台文档[EB/OL]. https://open.bigmodel.cn/, 2025.",
        "[6] DeepSeek. DeepSeek API 文档[EB/OL]. https://api-docs.deepseek.com/, 2025.",
        "[7] 阿里云. 通义千问 DashScope 兼容模式[EB/OL]. https://help.aliyun.com/zh/model-studio/, 2025.",
        "[8] 浙江理工大学信息学院. 《创新实践》指导书[Z]. 杭州: 浙江理工大学, 2017.",
    ]
    for r in refs:
        p = doc.add_paragraph()
        set_paragraph_format(p, first_indent=False)
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.74)
        add_text(p, r, size=10.5)

    h1(doc, "附录　实验主要函数与代码")
    h2(doc, "附录 A　查询主循环（节选）")
    pre(
        doc,
        "def ask(query, graph, llm, max_rounds=4):\n"
        "    evidence = Evidence()\n"
        "    queue = seed_subquestions(query)\n"
        "    rounds, draft, mark = 0, \"\", \"NEED\"\n"
        "    while rounds < max_rounds:\n"
        "        rounds += 1\n"
        "        sub = queue.pop(0) if queue else query\n"
        "        retrieve(sub, graph, evidence)\n"
        "        draft = clean_answer(\n"
        "            llm.predict(SYNTH_PROMPT.format(\n"
        "                query=query, evidence=evidence.as_prompt()))\n"
        "        )\n"
        "        mark = verify(query, draft, evidence, llm)\n"
        "        if mark in (\"OK\", \"UNKNOWN\"):\n"
        "            break\n"
        "        nxt = plan(query, evidence, llm, rounds, max_rounds)\n"
        "        if nxt.startswith(\"FOLLOWUP:\"):\n"
        "            queue.append(nxt.split(\":\", 1)[1].strip())\n"
        "        elif nxt == \"ANSWER\":\n"
        "            break\n"
        "    return {\"answer\": draft, \"path\": evidence.subquestions, ...}",
    )
    h2(doc, "附录 B　模板图查询（节选）")
    pre(
        doc,
        "def graph_query(self, name, hops=2, limit=30):\n"
        "    hops = max(1, min(int(hops), 3))\n"
        "    query = f\"\"\"\n"
        "    MATCH (start:Entity) WHERE start.name CONTAINS $name\n"
        "    MATCH path = (start)-[*1..{hops}]-(other:Entity)\n"
        "    WITH start, other, relationships(path) AS rels\n"
        "    UNWIND rels AS r\n"
        "    RETURN DISTINCT startNode(r).name AS src,\n"
        "           r.name AS rel, endNode(r).name AS dst\n"
        "    LIMIT $limit\n"
        "    \"\"\"",
    )
    body(
        doc,
        "完整源码见仓库 agents/、graphrag/graph.py、webdemo.py、eval/run_eval.py。密钥仅存在本地 .env，不入库。",
    )

    h1(doc, "附表　任务清单")
    caption_tab(doc, "表6 组员分工（提交前请与组员核对百分比）")
    add_table(
        doc,
        ["任务", "金科丞", "组员二（姓名：　　　）"],
        [
            ["选题、可行性分析报告、开题", "50%", "50%"],
            ["环境与图谱构建（Neo4j、前三章入库）", "60%", "40%"],
            ["智能体查询（规划 / 检索 / 成答 / 校验）", "50%", "50%"],
            ["多模型接入与网页（问答、子图、性能页）", "50%", "50%"],
            ["封闭评测、综合结论与实验分析", "50%", "50%"],
            ["课程设计报告、答辩 PPT 与演示", "50%", "50%"],
        ],
    )
    body(
        doc,
        "说明：任务书规定每组最多 2 人。上表默认均分实现与文档，建图环境由金科丞偏重。组员二姓名、学号及实际比例以小组最终确认为准。",
    )

    doc.save(OUT)
    print("wrote", OUT, "size", OUT.stat().st_size)


if __name__ == "__main__":
    build()
